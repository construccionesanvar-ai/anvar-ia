# Genera public/descargas/plantilla-roi-automatizacion.xlsx (requiere: pip install openpyxl)
#   python3 scripts/plantillas/roi.py
# openpyxl escribe las fórmulas sin valores calculados. Excel y LibreOffice las calculan al
# abrir, pero las vistas previas (Drive, WhatsApp, correo) muestran celdas vacías. Para dejar
# los valores guardados antes de publicar: abre el archivo en Excel o LibreOffice Calc,
# revisa que ninguna celda muestre #¡VALOR! o #¿NOMBRE? y guárdalo (Ctrl+S) como .xlsx.
# Celdas amarillas con texto azul = datos que ingresa el usuario. Negro = fórmulas.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
import sys

SALIDA = sys.argv[1] if len(sys.argv) > 1 else 'public/descargas/plantilla-roi-automatizacion.xlsx'
F = 'Arial'
AZUL = Font(name=F, size=11, color='0000FF', bold=True)
NEGRO = Font(name=F, size=11)
NEGRITA = Font(name=F, size=11, bold=True)
TITULO = Font(name=F, size=16, bold=True, color='14181B')
SUB = Font(name=F, size=10, color='4B5652')
SECCION = Font(name=F, size=11, bold=True, color='FFFFFF')
RES = Font(name=F, size=12, bold=True, color='14181B')
AMARILLO = PatternFill('solid', fgColor='FFFF00')
OSCURO = PatternFill('solid', fgColor='101A1E')
CLARO = PatternFill('solid', fgColor='F4F5F1')
AMBAR = PatternFill('solid', fgColor='F6E7C8')
linea = Side(style='thin', color='B4BAB0')
BORDE = Border(left=linea, right=linea, top=linea, bottom=linea)
PESOS = '$#,##0;($#,##0);-'
HORAS = '#,##0" h";(#,##0" h");-'
PCT = '0.0%;(0.0%);-'
MESES = '0.0" meses"'

wb = Workbook()

# ------------------------------------------------------------ Calculadora
ws = wb.active
ws.title = 'Calculadora'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 52
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 62

ws['A1'] = 'Plantilla de ROI de automatización'
ws['A1'].font = TITULO
ws['A2'] = 'ANVAR TECH · IA & Automatización · ia.anvartech.cl · Versión 1.0 (25/09/2026)'
ws['A2'].font = SUB
ws['A3'] = 'Ingresa tus datos en las celdas amarillas (texto azul). Todo lo demás se calcula solo. Los valores que trae son un ejemplo.'
ws['A3'].font = SUB

def seccion(fila, texto):
    for col in 'ABC':
        ws[f'{col}{fila}'].fill = OSCURO
    ws[f'A{fila}'] = texto
    ws[f'A{fila}'].font = SECCION
    ws[f'C{fila}'] = 'Nota'
    ws[f'C{fila}'].font = SECCION

def entrada(fila, etiqueta, valor, formato, nota):
    ws[f'A{fila}'] = etiqueta
    ws[f'A{fila}'].font = NEGRO
    c = ws[f'B{fila}']
    c.value = valor
    c.font = AZUL
    c.fill = AMARILLO
    c.number_format = formato
    c.border = BORDE
    c.alignment = Alignment(horizontal='right')
    ws[f'C{fila}'] = nota
    ws[f'C{fila}'].font = SUB
    ws[f'C{fila}'].alignment = Alignment(wrap_text=True, vertical='top')

def resultado(fila, etiqueta, formula, formato, nota, destacado=False):
    ws[f'A{fila}'] = etiqueta
    ws[f'A{fila}'].font = NEGRITA if destacado else NEGRO
    c = ws[f'B{fila}']
    c.value = formula
    c.font = RES if destacado else NEGRO
    c.number_format = formato
    c.border = BORDE
    c.alignment = Alignment(horizontal='right')
    if destacado:
        for col in 'AB':
            ws[f'{col}{fila}'].fill = AMBAR
    ws[f'C{fila}'] = nota
    ws[f'C{fila}'].font = SUB
    ws[f'C{fila}'].alignment = Alignment(wrap_text=True, vertical='top')

seccion(5, '1. El proceso hoy')
entrada(6, 'Personas que hacen la tarea', 5, '0', 'Cuántas personas hacen este mismo trabajo.')
entrada(7, 'Veces por semana que cada persona la hace', 10, '0', 'Frecuencia. Ej.: 2 veces al día, 5 días = 10.')
entrada(8, 'Minutos que toma cada vez', 36, '0" min"', 'Mídelo con un reloj en casos reales, incluido uno que salga mal.')
entrada(9, 'Semanas trabajadas al año', 44, '0', 'Supuesto: 52 semanas menos 3 de feriado legal (15 días hábiles, Código del Trabajo) y unas 5 de feriados, licencias y ausencias.')
entrada(10, 'Costo empresa por hora ($)', 9000, PESOS, 'Sueldo bruto más aportes del empleador, dividido por las horas trabajadas al mes. Lo tiene quien hace las remuneraciones.')

seccion(12, '2. La automatización')
entrada(13, 'Parte del tiempo que se automatiza', 0.6, PCT, 'Si no sabes, deja 60%. El resto queda para revisión humana y excepciones.')
entrada(14, 'Costo de implementación ($, neto)', 1640000, PESOS, 'Lo que cuesta dejarla funcionando. Ejemplo: piloto desde UF 40 + IVA (con UF de $41.000). Una Automatización Express parte en $199.900 + IVA.')
entrada(15, 'Costo mensual de mantención y licencias ($)', 0, PESOS, 'Soporte, suscripciones o licencias que necesite. 0 si no hay.')

seccion(17, '3. Resultados')
resultado(18, 'Horas al año que hoy toma el proceso', '=B6*B7*B8/60*B9', HORAS, 'Personas × veces por semana × minutos ÷ 60 × semanas.')
resultado(19, 'Costo anual del proceso hoy', '=B18*B10', PESOS, 'Horas al año × costo por hora.')
resultado(20, 'Horas que se liberan al año', '=B18*B13', HORAS, 'Horas al año × parte automatizable.', True)
resultado(21, 'Horas liberadas por persona a la semana', '=IF(B6*B9>0,B20/B6/B9,0)', '0.0" h"', 'Para explicarlo al equipo: cuánto tiempo recupera cada persona.')
resultado(22, 'Ahorro bruto anual', '=B20*B10', PESOS, 'Valor del tiempo liberado.')
resultado(23, 'Costo anual de mantención', '=B15*12', PESOS, 'Mantención mensual × 12.')
resultado(24, 'Ahorro neto anual', '=B22-B23', PESOS, 'Ahorro bruto menos mantención.', True)
resultado(25, 'Payback (meses para recuperar la inversión)', '=IF(B24>0,B14/(B24/12),"No se recupera")', MESES, 'Implementación ÷ ahorro neto mensual.', True)
resultado(26, 'ROI a 1 año', '=IF(B14>0,(B24-B14)/B14,0)', PCT, '(Ahorro neto del primer año − implementación) ÷ implementación.')
resultado(27, 'ROI a 3 años', '=IF(B14>0,(B24*3-B14)/B14,0)', PCT, '(Ahorro neto de tres años − implementación) ÷ implementación.')

ws['A29'] = 'Estimación referencial basada en los datos ingresados. El resultado real depende del proceso, la implementación y el contexto operacional. No incluye el costo de errores, reprocesos ni atrasos, que suele ser mayor que el de las horas.'
ws['A29'].font = SUB
ws['A29'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A29:C29')
ws.row_dimensions[29].height = 42

ws['A31'] = 'Leyenda'
ws['A31'].font = NEGRITA
ws['A32'] = 'Dato que ingresas'
ws['B32'].fill = AMARILLO
ws['B32'].value = 'azul'
ws['B32'].font = AZUL
ws['A33'] = 'Resultado principal'
ws['B33'].fill = AMBAR
ws['A34'] = 'Calculadora en línea y guías: https://ia.anvartech.cl/calculadora-roi-automatizacion'
ws['A34'].font = SUB
for f in range(32, 34):
    ws[f'A{f}'].font = NEGRO

dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='1', showErrorMessage=True, errorTitle='Porcentaje', error='Ingresa un porcentaje entre 0% y 100%.')
ws.add_data_validation(dv)
dv.add('B13')
dv2 = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True, errorTitle='Valor', error='Ingresa un número mayor o igual a 0.')
ws.add_data_validation(dv2)
for c in ['B6', 'B7', 'B8', 'B9', 'B10', 'B14', 'B15']:
    dv2.add(c)
ws['B9'].comment = Comment('Supuesto de ANVAR TECH, el mismo de la calculadora en línea. Cámbialo si tu equipo trabaja otras semanas.', 'ANVAR TECH')
ws.freeze_panes = 'A5'

# ------------------------------------------------------ Comparar procesos
cp = wb.create_sheet('Comparar procesos')
cp.sheet_view.showGridLines = False
cp['A1'] = 'Compara varios procesos para decidir por cuál partir'
cp['A1'].font = TITULO
cp['A2'] = 'Una fila por proceso. Usa las semanas y la columna de % de la hoja Calculadora como referencia. La fila 5 es un ejemplo: reemplázala.'
cp['A2'].font = SUB
enc = ['Proceso', 'Personas', 'Veces por semana', 'Minutos por vez', 'Costo por hora ($)', '% automatizable', 'Horas al año', 'Valor anual del tiempo', 'Horas liberadas al año', 'Valor liberado al año']
anchos = [34, 11, 16, 15, 18, 16, 14, 22, 20, 20]
for i, (t, w) in enumerate(zip(enc, anchos)):
    col = chr(65 + i)
    c = cp[f'{col}4']
    c.value = t
    c.font = SECCION
    c.fill = OSCURO
    c.alignment = Alignment(wrap_text=True, vertical='center')
    cp.column_dimensions[col].width = w
cp.row_dimensions[4].height = 32
ejemplo = ['Llenar formularios de despacho', 3, 15, 12, 8500, 0.7]
for fila in range(5, 15):
    for i in range(6):
        col = chr(65 + i)
        c = cp[f'{col}{fila}']
        c.fill = AMARILLO
        c.font = AZUL
        c.border = BORDE
        if fila == 5:
            c.value = ejemplo[i]
    cp[f'E{fila}'].number_format = PESOS
    cp[f'F{fila}'].number_format = PCT
    cp[f'G{fila}'] = f'=IF(A{fila}="","",B{fila}*C{fila}*D{fila}/60*Calculadora!$B$9)'
    cp[f'H{fila}'] = f'=IF(A{fila}="","",G{fila}*E{fila})'
    cp[f'I{fila}'] = f'=IF(A{fila}="","",G{fila}*F{fila})'
    cp[f'J{fila}'] = f'=IF(A{fila}="","",I{fila}*E{fila})'
    for col, fmt in (('G', HORAS), ('H', PESOS), ('I', HORAS), ('J', PESOS)):
        cp[f'{col}{fila}'].number_format = fmt
        cp[f'{col}{fila}'].font = NEGRO
        cp[f'{col}{fila}'].border = BORDE
cp['A16'] = 'Total'
cp['A16'].font = NEGRITA
for col, fmt in (('G', HORAS), ('H', PESOS), ('I', HORAS), ('J', PESOS)):
    cp[f'{col}16'] = f'=SUM({col}5:{col}14)'
    cp[f'{col}16'].number_format = fmt
    cp[f'{col}16'].font = NEGRITA
cp['A18'] = 'Parte por el proceso con más valor liberado al año que además se repita igual cada vez y tenga su forma de hacerse escrita.'
cp['A18'].font = SUB
cp.freeze_panes = 'A5'

# --------------------------------------------------------- Cómo usarla
gu = wb.create_sheet('Cómo usarla')
gu.sheet_view.showGridLines = False
gu.column_dimensions['A'].width = 110
lineas = [
    ('Cómo usar esta plantilla', TITULO),
    ('', NEGRO),
    ('1. En la hoja Calculadora, completa las celdas amarillas con los datos de un proceso real.', NEGRO),
    ('2. Mide la duración con reloj en al menos cinco casos reales, incluido uno que salga mal: ahí se esconde buena parte del tiempo.', NEGRO),
    ('3. Usa el costo empresa por hora, no el sueldo líquido. Si varias personas con distinto sueldo hacen la tarea, usa un promedio.', NEGRO),
    ('4. Si tienes varios procesos candidatos, compáralos en la hoja Comparar procesos y parte por el de más valor liberado.', NEGRO),
    ('', NEGRO),
    ('Cómo leer el resultado', NEGRITA),
    ('Payback: meses para recuperar la inversión con el ahorro neto. Bajo 12 meses suele ser una decisión fácil; sobre 36, conviene revisar si vale la pena.', NEGRO),
    ('ROI: cuánto se gana por cada peso invertido, descontada la inversión. 150% a un año significa recuperar la inversión y 1,5 veces más.', NEGRO),
    ('Horas liberadas: el tiempo vale como ahorro solo si se ocupa en algo útil. Si no, el beneficio real es menor.', NEGRO),
    ('', NEGRO),
    ('Qué no incluye', NEGRITA),
    ('El costo de errores, reprocesos y atrasos; la curva de aprendizaje de las primeras semanas; y los procesos nuevos que se vuelven posibles.', NEGRO),
    ('', NEGRO),
    ('Aviso', NEGRITA),
    ('Estimación referencial basada en los datos ingresados. El resultado real depende del proceso, la implementación y el contexto operacional.', NEGRO),
    ('', NEGRO),
    ('Hecha por ANVAR TECH (automatización de procesos, datos e IA aplicada a operaciones, Chile).', SUB),
    ('Calculadora en línea: https://ia.anvartech.cl/calculadora-roi-automatizacion', SUB),
    ('Guía de costos: https://ia.anvartech.cl/recursos/cuanto-cuesta-automatizar-proceso-chile', SUB),
    ('Puedes usarla y compartirla libremente. Contacto: contacto@anvartech.cl', SUB),
]
for i, (t, f) in enumerate(lineas, start=1):
    gu[f'A{i}'] = t
    gu[f'A{i}'].font = f
    gu[f'A{i}'].alignment = Alignment(wrap_text=True, vertical='top')

wb.properties.title = 'Plantilla de ROI de automatización'
wb.properties.creator = 'ANVAR TECH'
wb.properties.subject = 'Cálculo de ahorro, payback y ROI de automatizar un proceso'
wb.properties.keywords = 'ROI automatización, payback, ahorro, procesos'
wb.active = 0
wb.save(SALIDA)
print('ok', SALIDA)
